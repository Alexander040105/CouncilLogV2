import io
import uuid
from datetime import date, datetime, timezone

from fastapi import APIRouter, Depends, Response
from pydantic import BaseModel, Field
from sqlalchemy import func
from sqlmodel import select

from ..deps import Membership, Session, authorize
from ..errors import APIError, not_found
from ..models import (BudgetExpense, Document, ExpenseReceipt, FundTransaction,
                      Position, Profile, Project, ProjectBudget, SchoolYear)
from ..pagination import org_today
from ..services import storage
from ..services.audit import audit
from ..services.idempotent import add_deduped, deduped_response
from ..services.ratelimit import check_rate_limit

router = APIRouter(tags=["finance"])


# ── Helpers ───────────────────────────────────────────────────────────────

def _editable(recorded_by: uuid.UUID, day: date, member: Membership) -> bool:
    """Journal rule: creator on their own same-day record, or owner anytime."""
    return str(recorded_by) == member.user_id and day == org_today() \
        or member.role == "owner"


async def _signed_balance(session, org_id: uuid.UUID, through: date | None = None) -> int:
    """SUM(deposits) - SUM(withdrawals); `through` bounds by transacted_on."""
    q = (select(func.coalesce(func.sum(FundTransaction.amount_centavos), 0))
         .where(FundTransaction.org_id == org_id,
                FundTransaction.kind == "deposit"))
    w = (select(func.coalesce(func.sum(FundTransaction.amount_centavos), 0))
         .where(FundTransaction.org_id == org_id,
                FundTransaction.kind == "withdrawal"))
    if through is not None:
        q = q.where(FundTransaction.transacted_on <= through)
        w = w.where(FundTransaction.transacted_on <= through)
    dep = (await session.execute(q)).scalar_one()
    wit = (await session.execute(w)).scalar_one()
    return int(dep) - int(wit)


async def _spent(session, budget_id: uuid.UUID) -> int:
    return int((await session.execute(
        select(func.coalesce(func.sum(BudgetExpense.amount_centavos), 0))
        .where(BudgetExpense.budget_id == budget_id))).scalar_one())


async def _get_budget(session, org_id: uuid.UUID, budget_id: uuid.UUID) -> ProjectBudget:
    b = await session.get(ProjectBudget, budget_id)
    if b is None or b.org_id != org_id:
        raise not_found("budget")
    return b


async def _get_expense(session, org_id: uuid.UUID, expense_id: uuid.UUID) -> BudgetExpense:
    e = await session.get(BudgetExpense, expense_id)
    if e is None or e.org_id != org_id:
        raise not_found("expense")
    return e


async def _verify_upload(session, org_id: uuid.UUID, storage_path: str, mime: str) -> None:
    if not storage_path.startswith(f"{org_id}/"):
        raise APIError(422, "BAD_PATH", "File path not scoped to org")
    head = await storage.object_head(storage_path)
    if head is None:
        raise APIError(422, "UPLOAD_MISSING", "Uploaded object not found in storage")
    if not storage.check_magic_bytes(head, mime):
        raise APIError(422, "BAD_FILE_TYPE", "File content is not a valid receipt file")


async def _budget_payload(session, b: ProjectBudget, org_id: uuid.UUID) -> dict:
    spent = await _spent(session, b.id)
    txns = (await session.execute(select(FundTransaction).where(
        FundTransaction.budget_id == b.id))).scalars().all()
    expenses = (await session.execute(
        select(BudgetExpense).where(BudgetExpense.budget_id == b.id)
        .order_by(BudgetExpense.spent_on, BudgetExpense.created_at)
        .limit(2000))).scalars().all()
    receipts = (await session.execute(
        select(ExpenseReceipt).where(
            ExpenseReceipt.expense_id.in_([e.id for e in expenses])))
    ).scalars().all() if expenses else []
    rmap: dict[uuid.UUID, list] = {}
    for r in receipts:
        rmap.setdefault(r.expense_id, []).append(r)
    resolution_title = None
    if b.resolution_id:
        doc = await session.get(Document, b.resolution_id)
        resolution_title = doc.title if doc else None
    return {
        "budget": {**b.model_dump(),
                   "has_resolution_file": b.resolution_path is not None,
                   "resolution_title": resolution_title},
        "spent_centavos": spent,
        "remaining_centavos": b.allocated_centavos - spent,
        "withdrawal_txn": next((t for t in txns if t.kind == "withdrawal"), None),
        "return_txn": next((t for t in txns if t.kind == "deposit"), None),
        "expenses": [{**e.model_dump(),
                      "receipts": [r.model_dump() for r in rmap.get(e.id, [])]}
                     for e in expenses],
    }


# ── Bankbook ──────────────────────────────────────────────────────────────
class TxnIn(BaseModel):
    kind: str = Field(pattern="^(deposit|withdrawal)$")
    amount_centavos: int = Field(gt=0)
    transacted_on: date | None = None
    source_label: str | None = Field(default=None, max_length=60)
    note: str | None = Field(default=None, max_length=500)


@router.get("/orgs/{org_id}/fund")
async def get_fund(org_id: uuid.UUID, session: Session,
                   member: Membership = Depends(authorize())):
    balance = await _signed_balance(session, org_id)
    rows = (await session.execute(
        select(FundTransaction, Project.title)
        .join(ProjectBudget, ProjectBudget.id == FundTransaction.budget_id, isouter=True)
        .join(Project, Project.id == ProjectBudget.project_id, isouter=True)
        .where(FundTransaction.org_id == org_id)
        .order_by(FundTransaction.transacted_on.desc(),
                  FundTransaction.created_at.desc()).limit(500))).all()
    return {"balance_centavos": balance,
            "transactions": [{**t.model_dump(), "project_title": title}
                             for t, title in rows]}


@router.post("/orgs/{org_id}/fund/transactions", status_code=201)
async def create_txn(org_id: uuid.UUID, body: TxnIn, session: Session,
                     member: Membership = Depends(authorize("officer"))):
    t = FundTransaction(org_id=org_id, kind=body.kind,
                        amount_centavos=body.amount_centavos,
                        transacted_on=body.transacted_on or org_today(),
                        source_label=body.source_label, note=body.note,
                        created_by=uuid.UUID(member.user_id))
    session.add(t)
    await audit(session, org_id=org_id, actor_id=member.user_id,
                action="fund.transaction", entity_type="fund_transaction",
                entity_id=t.id,
                metadata={"kind": t.kind, "amount_centavos": t.amount_centavos})
    await session.commit()
    return {"data": t}


class TxnPatch(BaseModel):
    amount_centavos: int | None = Field(default=None, gt=0)
    transacted_on: date | None = None
    source_label: str | None = None   # tri-state
    note: str | None = None           # tri-state


def _guarded_txn(t: FundTransaction, member: Membership) -> None:
    if t.budget_id is not None:
        raise APIError(422, "LINKED_TRANSACTION",
                       "This entry belongs to a project budget — change it through the budget")
    if not _editable(t.created_by, t.transacted_on, member):
        raise APIError(403, "FORBIDDEN",
                       "Only the creator same-day, or an owner, can edit transactions")


@router.patch("/orgs/{org_id}/fund/transactions/{txn_id}")
async def patch_txn(org_id: uuid.UUID, txn_id: uuid.UUID, body: TxnPatch,
                    session: Session, member: Membership = Depends(authorize("officer"))):
    t = await session.get(FundTransaction, txn_id)
    if t is None or t.org_id != org_id:
        raise not_found("transaction")
    _guarded_txn(t, member)
    for k, v in body.model_dump(exclude_unset=True).items():
        setattr(t, k, v)
    await audit(session, org_id=org_id, actor_id=member.user_id,
                action="fund.transaction_edited", entity_type="fund_transaction",
                entity_id=t.id)
    await session.commit()
    return {"data": t}


@router.delete("/orgs/{org_id}/fund/transactions/{txn_id}")
async def delete_txn(org_id: uuid.UUID, txn_id: uuid.UUID, session: Session,
                     member: Membership = Depends(authorize("officer"))):
    t = await session.get(FundTransaction, txn_id)
    if t is None or t.org_id != org_id:
        raise not_found("transaction")
    _guarded_txn(t, member)
    await audit(session, org_id=org_id, actor_id=member.user_id,
                action="fund.transaction_deleted", entity_type="fund_transaction",
                entity_id=t.id,
                metadata={"kind": t.kind, "amount_centavos": t.amount_centavos})
    await session.delete(t)
    await session.commit()
    return {"ok": True}


# ── Resolution upload ─────────────────────────────────────────────────────
class FileSign(BaseModel):
    mime: str
    byte_size: int = Field(gt=0, le=5 * 1024 * 1024)


@router.post("/orgs/{org_id}/fund/resolution-sign", status_code=201)
async def sign_resolution(org_id: uuid.UUID, body: FileSign, session: Session,
                          member: Membership = Depends(authorize("officer"))):
    """Mint an upload URL for a resolution scan — signed BEFORE the budget
    exists; the file is uploaded, then referenced at budget create."""
    await check_rate_limit(session, f"receiptsign:{member.user_id}",
                           limit=60, window_seconds=3600)
    storage.validate_upload_declared(body.mime, body.byte_size)
    path = f"{org_id}/resolutions/{uuid.uuid4()}"
    return {"path": path, "upload_url": await storage.signed_upload_url(path)}


class UploadedFile(BaseModel):
    storage_path: str
    mime: str
    byte_size: int = Field(gt=0, le=5 * 1024 * 1024)


# ── Budgets ───────────────────────────────────────────────────────────────
class BudgetIn(BaseModel):
    allocated_centavos: int = Field(gt=0)
    resolution_id: uuid.UUID | None = None
    resolution: UploadedFile | None = None
    source_label: str | None = Field(default=None, max_length=60)
    note: str | None = Field(default=None, max_length=500)


async def _check_resolution_doc(session, org_id: uuid.UUID, doc_id: uuid.UUID) -> Document:
    d = await session.get(Document, doc_id)
    if d is None or d.org_id != org_id:
        raise not_found("document")
    if d.doc_type != "board_resolution":
        raise APIError(422, "WRONG_DOC_TYPE",
                       "That document isn't a board resolution — pick one with type board_resolution")
    return d


@router.post("/orgs/{org_id}/projects/{project_id}/budget", status_code=201)
async def create_budget(org_id: uuid.UUID, project_id: uuid.UUID, body: BudgetIn,
                        session: Session, member: Membership = Depends(authorize("officer"))):
    p = await session.get(Project, project_id)
    if p is None or p.org_id != org_id:
        raise not_found("project")
    existing = (await session.execute(select(func.count()).where(
        ProjectBudget.project_id == project_id))).scalar_one()
    if existing:
        raise APIError(409, "BUDGET_EXISTS", "This project already has a budget")
    if body.resolution_id:
        await _check_resolution_doc(session, org_id, body.resolution_id)
    resolution_path = resolution_mime = None
    resolution_size = None
    if body.resolution:
        await _verify_upload(session, org_id, body.resolution.storage_path,
                             body.resolution.mime)
        resolution_path = body.resolution.storage_path
        resolution_mime = body.resolution.mime
        resolution_size = body.resolution.byte_size
    b = ProjectBudget(org_id=org_id, project_id=project_id,
                      allocated_centavos=body.allocated_centavos,
                      source_label=body.source_label,
                      resolution_id=body.resolution_id,
                      resolution_path=resolution_path,
                      resolution_mime=resolution_mime,
                      resolution_byte_size=resolution_size,
                      note=body.note, created_by=uuid.UUID(member.user_id))
    session.add(b)
    await session.flush()
    txn = FundTransaction(org_id=org_id, kind="withdrawal",
                          amount_centavos=b.allocated_centavos,
                          transacted_on=org_today(),
                          source_label=body.source_label,
                          note=f'Withdrawal for "{p.title}"', budget_id=b.id,
                          created_by=uuid.UUID(member.user_id))
    session.add(txn)
    await audit(session, org_id=org_id, actor_id=member.user_id,
                action="budget.created", entity_type="project_budget",
                entity_id=b.id,
                metadata={"project_id": str(project_id),
                          "allocated_centavos": b.allocated_centavos})
    await session.commit()
    return {"data": b}


@router.get("/orgs/{org_id}/projects/{project_id}/budget")
async def get_project_budget(org_id: uuid.UUID, project_id: uuid.UUID,
                             session: Session, member: Membership = Depends(authorize())):
    b = (await session.execute(select(ProjectBudget).where(
        ProjectBudget.project_id == project_id))).scalars().first()
    if b is None or b.org_id != org_id:
        raise not_found("budget")
    return await _budget_payload(session, b, org_id)


@router.get("/orgs/{org_id}/budgets")
async def list_budgets(org_id: uuid.UUID, session: Session,
                       member: Membership = Depends(authorize())):
    budgets = (await session.execute(
        select(ProjectBudget, Project.title)
        .join(Project, Project.id == ProjectBudget.project_id)
        .where(ProjectBudget.org_id == org_id)
        .order_by(Project.target_date.desc().nulls_last())
        .limit(500))).all()
    if not budgets:
        return {"data": []}
    spent_rows = (await session.execute(
        select(BudgetExpense.budget_id,
               func.coalesce(func.sum(BudgetExpense.amount_centavos), 0))
        .where(BudgetExpense.budget_id.in_([b.id for b, _ in budgets]))
        .group_by(BudgetExpense.budget_id))).all()
    smap = {bid: int(s) for bid, s in spent_rows}
    return {"data": [{**b.model_dump(), "project_title": title,
                      "spent_centavos": smap.get(b.id, 0),
                      "remaining_centavos": b.allocated_centavos - smap.get(b.id, 0)}
                     for b, title in budgets]}


class BudgetPatch(BaseModel):
    allocated_centavos: int | None = Field(default=None, gt=0)
    source_label: str | None = None
    note: str | None = None
    resolution_id: uuid.UUID | None = None
    resolution: UploadedFile | None = None  # explicit null clears the file


@router.patch("/orgs/{org_id}/budgets/{budget_id}")
async def patch_budget(org_id: uuid.UUID, budget_id: uuid.UUID, body: BudgetPatch,
                       session: Session, member: Membership = Depends(authorize("officer"))):
    b = await _get_budget(session, org_id, budget_id)
    if b.status == "closed":
        raise APIError(422, "BUDGET_CLOSED",
                       "This budget is closed — its numbers are locked in the bankbook")
    old_resolution_path = None
    if body.allocated_centavos is not None and body.allocated_centavos != b.allocated_centavos:
        b.allocated_centavos = body.allocated_centavos
        txn = (await session.execute(select(FundTransaction).where(
            FundTransaction.budget_id == b.id,
            FundTransaction.kind == "withdrawal"))).scalars().first()
        if txn is not None:
            txn.amount_centavos = body.allocated_centavos
    if "source_label" in body.model_fields_set:
        b.source_label = body.source_label
    if "note" in body.model_fields_set:
        b.note = body.note
    if "resolution_id" in body.model_fields_set:
        if body.resolution_id is not None:
            await _check_resolution_doc(session, org_id, body.resolution_id)
        b.resolution_id = body.resolution_id
    if "resolution" in body.model_fields_set:
        old_resolution_path = b.resolution_path
        if body.resolution is None:
            b.resolution_path = b.resolution_mime = b.resolution_byte_size = None
        else:
            await _verify_upload(session, org_id, body.resolution.storage_path,
                                 body.resolution.mime)
            b.resolution_path = body.resolution.storage_path
            b.resolution_mime = body.resolution.mime
            b.resolution_byte_size = body.resolution.byte_size
    await audit(session, org_id=org_id, actor_id=member.user_id,
                action="budget.updated", entity_type="project_budget", entity_id=b.id)
    await session.commit()
    if old_resolution_path and old_resolution_path != b.resolution_path:
        await storage.delete_object(old_resolution_path)
    return {"data": b}


@router.get("/orgs/{org_id}/budgets/{budget_id}/resolution/url")
async def resolution_url(org_id: uuid.UUID, budget_id: uuid.UUID,
                         session: Session, member: Membership = Depends(authorize())):
    b = await _get_budget(session, org_id, budget_id)
    if not b.resolution_path:
        raise APIError(404, "NO_RESOLUTION_FILE",
                       "No resolution file is attached to this budget")
    return {"download_url": await storage.signed_download_url(b.resolution_path)}


@router.post("/orgs/{org_id}/budgets/{budget_id}/close")
async def close_budget(org_id: uuid.UUID, budget_id: uuid.UUID, session: Session,
                       member: Membership = Depends(authorize("officer"))):
    b = await _get_budget(session, org_id, budget_id)
    if b.status == "closed":
        raise APIError(422, "BUDGET_CLOSED", "This budget is already closed")
    p = await session.get(Project, b.project_id)
    remaining = b.allocated_centavos - await _spent(session, b.id)
    if remaining < 0:
        raise APIError(422, "OVERSPENT",
                       "Spending is over the allocated amount — raise the allocation first")
    txn = None
    if remaining > 0:
        txn = FundTransaction(org_id=org_id, kind="deposit",
                              amount_centavos=remaining,
                              transacted_on=org_today(),
                              source_label=b.source_label,
                              note=f'Returned remaining fund — "{p.title}"',
                              budget_id=b.id,
                              created_by=uuid.UUID(member.user_id))
        session.add(txn)
        await session.flush()
    b.status = "closed"
    b.closed_at = datetime.now(timezone.utc)
    await audit(session, org_id=org_id, actor_id=member.user_id,
                action="budget.closed", entity_type="project_budget",
                entity_id=b.id,
                metadata={"remaining_centavos": remaining})
    await session.commit()
    return {"data": b, "return_txn": txn,
            "balance_centavos": await _signed_balance(session, org_id)}


@router.delete("/orgs/{org_id}/budgets/{budget_id}")
async def delete_budget(org_id: uuid.UUID, budget_id: uuid.UUID, session: Session,
                        member: Membership = Depends(authorize("officer"))):
    b = await _get_budget(session, org_id, budget_id)
    n_exp = (await session.execute(select(func.count()).where(
        BudgetExpense.budget_id == b.id))).scalar_one()
    if n_exp:
        raise APIError(422, "HAS_EXPENSES",
                       "This budget has expenses — close it instead of deleting")
    txns = (await session.execute(select(FundTransaction).where(
        FundTransaction.budget_id == b.id))).scalars().all()
    for t in txns:
        await session.delete(t)
    path = b.resolution_path
    await audit(session, org_id=org_id, actor_id=member.user_id,
                action="budget.deleted", entity_type="project_budget", entity_id=b.id)
    await session.delete(b)
    await session.commit()
    if path:
        await storage.delete_object(path)
    return {"ok": True}


# ── Expenses ──────────────────────────────────────────────────────────────
class ExpenseIn(BaseModel):
    vendor: str | None = Field(default=None, max_length=200)
    item: str = Field(min_length=1, max_length=300)
    amount_centavos: int = Field(gt=0)
    spent_on: date | None = None
    note: str | None = Field(default=None, max_length=500)
    client_request_id: str | None = Field(default=None, max_length=64)


@router.post("/orgs/{org_id}/budgets/{budget_id}/expenses", status_code=201)
async def create_expense(org_id: uuid.UUID, budget_id: uuid.UUID, body: ExpenseIn,
                         session: Session, member: Membership = Depends(authorize("officer"))):
    b = await _get_budget(session, org_id, budget_id)
    if b.status == "closed":
        raise APIError(422, "BUDGET_CLOSED",
                       "This budget is closed — reopening isn't supported")
    e = BudgetExpense(org_id=org_id, budget_id=b.id, vendor=body.vendor,
                      item=body.item, amount_centavos=body.amount_centavos,
                      spent_on=body.spent_on or org_today(), note=body.note,
                      recorded_by=uuid.UUID(member.user_id),
                      client_request_id=body.client_request_id)
    e, deduped = await add_deduped(session, e, BudgetExpense, org_id=org_id,
                                 client_request_id=body.client_request_id)
    if deduped:
        return deduped_response(e)
    await audit(session, org_id=org_id, actor_id=member.user_id,
                action="expense.created", entity_type="budget_expense",
                entity_id=e.id,
                metadata={"budget_id": str(b.id),
                          "amount_centavos": e.amount_centavos})
    await session.commit()
    return {"data": e}


class ExpensePatch(BaseModel):
    vendor: str | None = None
    item: str | None = Field(default=None, min_length=1, max_length=300)
    amount_centavos: int | None = Field(default=None, gt=0)
    spent_on: date | None = None
    note: str | None = None


@router.patch("/orgs/{org_id}/expenses/{expense_id}")
async def patch_expense(org_id: uuid.UUID, expense_id: uuid.UUID, body: ExpensePatch,
                        session: Session, member: Membership = Depends(authorize())):
    e = await _get_expense(session, org_id, expense_id)
    b = await _get_budget(session, org_id, e.budget_id)
    if b.status == "closed":
        raise APIError(422, "BUDGET_CLOSED", "This budget is closed")
    if not _editable(e.recorded_by, e.spent_on, member):
        raise APIError(403, "FORBIDDEN",
                       "Only the recorder same-day, or an owner, can edit expenses")
    for k, v in body.model_dump(exclude_unset=True).items():
        setattr(e, k, v)
    e.updated_at = datetime.now(timezone.utc)
    await audit(session, org_id=org_id, actor_id=member.user_id,
                action="expense.edited", entity_type="budget_expense", entity_id=e.id)
    await session.commit()
    return {"data": e}


@router.delete("/orgs/{org_id}/expenses/{expense_id}")
async def delete_expense(org_id: uuid.UUID, expense_id: uuid.UUID, session: Session,
                         member: Membership = Depends(authorize())):
    e = await _get_expense(session, org_id, expense_id)
    b = await _get_budget(session, org_id, e.budget_id)
    if b.status == "closed":
        raise APIError(422, "BUDGET_CLOSED", "This budget is closed")
    if not _editable(e.recorded_by, e.spent_on, member):
        raise APIError(403, "FORBIDDEN",
                       "Only the recorder same-day, or an owner, can delete expenses")
    receipts = (await session.execute(select(ExpenseReceipt).where(
        ExpenseReceipt.expense_id == e.id))).scalars().all()
    paths = [r.storage_path for r in receipts]
    await audit(session, org_id=org_id, actor_id=member.user_id,
                action="expense.deleted", entity_type="budget_expense",
                entity_id=e.id,
                metadata={"item": e.item, "amount_centavos": e.amount_centavos})
    await session.delete(e)
    await session.commit()
    for path in paths:
        await storage.delete_object(path)
    return {"ok": True}


# ── Receipts ──────────────────────────────────────────────────────────────
@router.post("/orgs/{org_id}/expenses/{expense_id}/receipts/sign", status_code=201)
async def sign_receipt(org_id: uuid.UUID, expense_id: uuid.UUID, body: FileSign,
                       session: Session, member: Membership = Depends(authorize("officer"))):
    await _get_expense(session, org_id, expense_id)
    await check_rate_limit(session, f"receiptsign:{member.user_id}",
                           limit=60, window_seconds=3600)
    storage.validate_upload_declared(body.mime, body.byte_size)
    path = f"{org_id}/receipts/{expense_id}/{uuid.uuid4()}"
    return {"path": path, "upload_url": await storage.signed_upload_url(path)}


@router.post("/orgs/{org_id}/expenses/{expense_id}/receipts", status_code=201)
async def attach_receipt(org_id: uuid.UUID, expense_id: uuid.UUID,
                         body: UploadedFile, session: Session,
                         member: Membership = Depends(authorize("officer"))):
    e = await _get_expense(session, org_id, expense_id)
    b = await _get_budget(session, org_id, e.budget_id)
    if b.status == "closed":
        raise APIError(422, "BUDGET_CLOSED", "This budget is closed")
    await _verify_upload(session, org_id, body.storage_path, body.mime)
    r = ExpenseReceipt(org_id=org_id, expense_id=e.id,
                       storage_path=body.storage_path, mime=body.mime,
                       byte_size=body.byte_size)
    session.add(r)
    await session.commit()
    return {"data": r}


@router.delete("/orgs/{org_id}/expenses/{expense_id}/receipts/{receipt_id}")
async def delete_receipt(org_id: uuid.UUID, expense_id: uuid.UUID,
                         receipt_id: uuid.UUID, session: Session,
                         member: Membership = Depends(authorize())):
    e = await _get_expense(session, org_id, expense_id)
    if not _editable(e.recorded_by, e.spent_on, member):
        raise APIError(403, "FORBIDDEN",
                       "Only the recorder same-day, or an owner, can remove receipts")
    r = await session.get(ExpenseReceipt, receipt_id)
    if r is None or r.expense_id != e.id:
        raise not_found("receipt")
    path = r.storage_path
    await session.delete(r)
    await session.commit()
    await storage.delete_object(path)
    return {"ok": True}


@router.get("/orgs/{org_id}/receipts/{receipt_id}/url")
async def receipt_url(org_id: uuid.UUID, receipt_id: uuid.UUID, session: Session,
                      member: Membership = Depends(authorize())):
    r = await session.get(ExpenseReceipt, receipt_id)
    if r is None or r.org_id != org_id:
        raise not_found("receipt")
    return {"download_url": await storage.signed_download_url(r.storage_path)}


# ── Financial report ──────────────────────────────────────────────────────
async def _report_data(session, org_id: uuid.UUID, project_id: uuid.UUID) -> dict:
    p = await session.get(Project, project_id)
    if p is None or p.org_id != org_id:
        raise not_found("project")
    b = (await session.execute(select(ProjectBudget).where(
        ProjectBudget.project_id == project_id))).scalars().first()
    if b is None or b.org_id != org_id:
        raise not_found("budget")
    payload = await _budget_payload(session, b, org_id)

    groups: dict[str | None, list] = {}
    for e in payload["expenses"]:
        groups.setdefault(e["vendor"], []).append(e)
    by_vendor = [{"vendor": v,
                  "items": [{"expense_id": i["id"], "item": i["item"],
                             "amount_centavos": i["amount_centavos"],
                             "spent_on": i["spent_on"],
                             "receipts": [{"id": r["id"], "mime": r["mime"]}
                                          for r in i["receipts"]]}
                            for i in items],
                  "subtotal_centavos": sum(i["amount_centavos"] for i in items)}
                 for v, items in groups.items()]

    ret = payload["return_txn"]
    return_block = None
    if ret is not None:
        return_block = {"deposited_on": ret.transacted_on,
                        "amount_centavos": ret.amount_centavos,
                        "balance_after_centavos": await _signed_balance(
                            session, org_id, through=ret.transacted_on)}

    sy = (await session.execute(select(SchoolYear).where(
        SchoolYear.org_id == org_id,
        SchoolYear.is_current == True))).scalars().first()  # noqa: E712
    signatories = {"auditor": None, "treasurer": None,
                   "president": None, "adviser": None}
    if sy is not None:
        positions = (await session.execute(select(Position).where(
            Position.org_id == org_id, Position.school_year_id == sy.id,
            Position.holder.is_not(None)))).scalars().all()
        holders = {pos.holder for pos in positions}
        profiles = {pr.id: pr.display_name for pr in (await session.execute(
            select(Profile).where(Profile.id.in_(holders)))).scalars().all()} \
            if holders else {}
        for pos in positions:
            t = (pos.title or "").lower()
            for key in signatories:
                if key in t and signatories[key] is None:
                    signatories[key] = profiles.get(pos.holder)

    return {
        "project": {"title": p.title, "event_type": p.event_type,
                    "target_date": p.target_date},
        "fund": {"source_label": b.source_label,
                 "allocated_centavos": b.allocated_centavos,
                 "withdrawn_on": payload["withdrawal_txn"].transacted_on
                 if payload["withdrawal_txn"] else None},
        "resolution": {"document_id": b.resolution_id,
                       "document_title": (await _budget_payload_title(
                           session, b.resolution_id)),
                       "has_file": b.resolution_path is not None,
                       "budget_id": b.id},
        "expenses_by_vendor": by_vendor,
        "totals": {"spent_centavos": payload["spent_centavos"],
                   "remaining_centavos": payload["remaining_centavos"]},
        "return": return_block,
        "signatories": signatories,
    }


async def _budget_payload_title(session, doc_id: uuid.UUID | None) -> str | None:
    if doc_id is None:
        return None
    d = await session.get(Document, doc_id)
    return d.title if d else None


@router.get("/orgs/{org_id}/projects/{project_id}/financial-report")
async def financial_report(org_id: uuid.UUID, project_id: uuid.UUID,
                           session: Session, member: Membership = Depends(authorize())):
    return await _report_data(session, org_id, project_id)


def _peso(centavos: int) -> str:
    return f"₱{centavos / 100:,.2f}"


@router.get("/orgs/{org_id}/projects/{project_id}/financial-report.docx")
async def financial_report_docx(org_id: uuid.UUID, project_id: uuid.UUID,
                                session: Session, member: Membership = Depends(authorize())):
    import docx

    data = await _report_data(session, org_id, project_id)
    d = docx.Document()
    d.add_heading("Financial Report Form (FRF)", level=0)
    proj = data["project"]
    d.add_heading("Activity Details", level=1)
    t = d.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    t.rows[0].cells[0].text = f"Title of Activity: {proj['title']}"
    t.rows[0].cells[1].text = f"Date: {proj['target_date'] or '—'}"
    t.rows[1].cells[0].text = f"Nature of Activity: {proj['event_type'] or '—'}"
    t.rows[1].cells[1].text = ""

    d.add_heading("Collection and Expenses", level=1)
    fund = data["fund"]
    d.add_paragraph(f"Source of Fund: {fund['source_label'] or 'Bankbook'}"
                    + (f" as of {fund['withdrawn_on']}" if fund["withdrawn_on"] else ""))
    t = d.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    t.rows[0].cells[0].text = "Total Fund:"
    t.rows[0].cells[1].text = _peso(fund["allocated_centavos"])
    t.rows[1].cells[0].text = "Less Expenses"
    t.rows[1].cells[1].text = ""

    for group in data["expenses_by_vendor"]:
        label = group["vendor"] or "Expenses"
        d.add_paragraph(label, style="List Number")
        for i in group["items"]:
            d.add_paragraph(f"{i['item']} — {_peso(i['amount_centavos'])}",
                            style="List Bullet")
        d.add_paragraph(f"Total: {_peso(group['subtotal_centavos'])}")

    totals = data["totals"]
    t = d.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    t.rows[0].cells[0].text = "Total Expenses"
    t.rows[0].cells[1].text = _peso(totals["spent_centavos"])
    t.rows[1].cells[0].text = "Total Remaining Fund"
    t.rows[1].cells[1].text = _peso(totals["remaining_centavos"])

    d.add_paragraph("Note: Please see attached pictures of receipts for proof of payment.")
    if data["return"]:
        r = data["return"]
        d.add_paragraph(
            f"Note: The bankbook reflects a deposit of {_peso(r['amount_centavos'])} "
            f"on {r['deposited_on']} — the remaining fund from "
            f"\"{proj['title']}\". The total bank balance after the deposit is "
            f"{_peso(r['balance_after_centavos'])}.")

    sig = data["signatories"]
    d.add_paragraph("")
    for label, key in (("AUDITED AND PREPARED BY:", "auditor"),
                       ("NOTED BY:", "treasurer"),
                       ("RECOMMENDING APPROVAL BY:", "president"),
                       ("APPROVED BY:", "adviser")):
        d.add_paragraph(label)
        d.add_paragraph(sig[key] or "____________________")

    buf = io.BytesIO()
    d.save(buf)
    slug = "".join(c if c.isalnum() else "-" for c in proj["title"].lower())[:60]
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition":
                 f'attachment; filename="{slug}-financial-report.docx"'})
